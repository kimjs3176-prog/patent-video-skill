"""Write copies of the source Excel lists with a last column naming each row's video file(s).

Usage: python -I tools/annotate_lists.py [--template A..I|S] <list.xlsx> [...]
The column lists the deliverable patents/{출원번호}.mp4 (see tools/export.py); --template X lists that template's
renders_* file name instead.
Copies are saved next to the originals as <name>_영상파일명.xlsx; originals are never modified.
The header row of each sheet is the row just above the first row that mentions a patent;
the new column goes right after the sheet's last used column and copies its style.
"""
import copy
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PAT = re.compile(r"(?<![0-9])(10-\d{4}-\d{7})(?![0-9])")

args = sys.argv[1:]
tpl = None
if "--template" in args:
    tpl = args[args.index("--template") + 1].upper()
    args = [a for i, a in enumerate(args) if a != "--template" and (i == 0 or args[i - 1] != "--template")]
master = json.load(open(os.path.join(ROOT, "data", "master.json"), encoding="utf-8"))
if tpl:
    suffix = "" if tpl == "A" else f"_{tpl.lower()}"
    SLUG = {m["appno"]: f'{m["order"]:03d}-{m["appno"]}{suffix}.mp4' for m in master}
else:
    SLUG = {m["appno"]: f'patents/{m["appno"]}.mp4' for m in master}


def files_in(texts):
    out = []
    for t in texts:
        t = str(t or "")
        for m in PAT.finditer(t):
            if "제외" in t[max(0, m.start() - 8):m.start()]:
                continue
            f = SLUG.get(m.group(1).replace("-", ""))
            if f and f not in out:
                out.append(f)
    return out


for path in args:
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        used = [(c.row, c.column) for row in ws.iter_rows() for c in row if c.value not in (None, "")]
        if not used:
            continue
        last_row = max(r for r, _ in used)
        last_col = max(c for _, c in used)
        first = next((r for r in range(1, last_row + 1) if files_in(ws.cell(r, c).value for c in range(1, last_col + 1))), None)
        if first is None:
            continue
        header = max(1, first - 1)
        col = last_col + 1
        h = ws.cell(header, col, "기술소개 영상 파일")
        h._style = copy.copy(ws.cell(header, last_col)._style)
        n = 0
        for r in range(first, last_row + 1):
            files = files_in(ws.cell(r, c).value for c in range(1, last_col + 1))
            if files:
                c = ws.cell(r, col, "\n".join(files))
                c._style = copy.copy(ws.cell(r, last_col)._style)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                n += 1
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 30
        print(os.path.basename(path), ws.title, "->", n, "rows")
    out = path[: path.rfind(".")] + "_영상파일명.xlsx"
    wb.save(out)
    print("saved", out)
